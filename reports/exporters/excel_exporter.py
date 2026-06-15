from __future__ import annotations

import io
from typing import Any

from core.logging import get_logger
from core.result import Result

logger = get_logger(__name__)


class ExcelExporter:
    async def export(
        self, data: list[dict[str, Any]], output_path: str = "", sheet_name: str = "Sheet1"
    ) -> Result[str]:
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        if data:
            headers = list(data[0].keys())
            ws.append(headers)
            for row in data:
                ws.append([row.get(h, "") for h in headers])
        if output_path:
            wb.save(output_path)
            return Result.ok(output_path)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Result.ok(buf.read().decode("latin1"))

    async def export_to_bytes(self, data: list[dict[str, Any]], sheet_name: str = "Sheet1") -> Result[bytes]:
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        if data:
            headers = list(data[0].keys())
            ws.append(headers)
            for row in data:
                ws.append([row.get(h, "") for h in headers])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Result.ok(buf.getvalue())

    async def export_multi_sheet(self, sheets: dict[str, list[dict[str, Any]]], output_path: str = "") -> Result[str]:
        import openpyxl

        wb = openpyxl.Workbook()
        first = True
        for sheet_name, rows in sheets.items():
            if first:
                ws = wb.active
                ws.title = sheet_name[:31]
                first = False
            else:
                ws = wb.create_sheet(title=sheet_name[:31])
            if rows:
                headers = list(rows[0].keys())
                ws.append(headers)
                for row in rows:
                    ws.append([row.get(h, "") for h in headers])
        if output_path:
            wb.save(output_path)
            return Result.ok(output_path)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Result.ok(buf.read().decode("latin1"))
