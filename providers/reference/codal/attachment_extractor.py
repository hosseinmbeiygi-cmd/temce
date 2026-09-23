from __future__ import annotations

from core.logging import get_logger
from core.result import Result

logger = get_logger(__name__)


class AttachmentExtractor:
    def __init__(self) -> None:
        self.supported_types = [".pdf", ".xlsx", ".xls", ".html", ".xml"]

    async def extract(self, content: bytes, file_type: str) -> Result[str]:
        file_type = file_type.lower()
        if file_type == ".pdf":
            return await self._extract_pdf(content)
        elif file_type in (".xlsx", ".xls"):
            return await self._extract_excel(content)
        elif file_type == ".html":
            return await self._extract_html(content)
        elif file_type == ".xml":
            return await self._extract_xml(content)
        return Result.fail(f"Unsupported file type: {file_type}")

    async def _extract_pdf(self, content: bytes) -> Result[str]:
        try:
            import PyPDF2

            reader = PyPDF2.PdfReader(content)
            text = "\n".join(page.extract_text() for page in reader.pages)
            return Result.ok(text)
        except ImportError:
            logger.warning("PyPDF2 not installed")
            return Result.ok("PDF extraction requires PyPDF2")
        except Exception as e:
            return Result.fail(f"PDF extraction failed: {e}")

    async def _extract_excel(self, content: bytes) -> Result[str]:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(content, read_only=True, data_only=True)
            rows: list[str] = []
            for ws in wb.worksheets:
                # Audit fix (C7): openpyxl's method is ``iter_rows`` —
                # ``iter_row`` never existed and would raise AttributeError.
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join(str(c) if c else "" for c in row))
            wb.close()
            return Result.ok("\n".join(rows))
        except ImportError:
            return Result.ok("Excel extraction requires openpyxl")
        except Exception as e:
            return Result.fail(f"Excel extraction failed: {e}")

    async def _extract_html(self, content: bytes) -> Result[str]:
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(content, "html.parser")
            return Result.ok(soup.get_text(separator="\n"))
        except ImportError:
            return Result.ok("HTML extraction requires beautifulsoup4")
        except Exception as e:
            return Result.fail(f"HTML extraction failed: {e}")

    async def _extract_xml(self, content: bytes) -> Result[str]:
        try:
            import xml.etree.ElementTree as ET

            root = ET.fromstring(content)
            return Result.ok(ET.tostring(root, encoding="unicode"))
        except Exception as e:
            return Result.fail(f"XML extraction failed: {e}")
