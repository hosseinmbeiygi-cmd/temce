"""Test parsing Codal Excel files with openpyxl."""

import io
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl

base = "codal_excel_files/آباد"
files = sorted(os.listdir(base))
if not files:
    print("No files found")
    sys.exit(1)

fp = os.path.join(base, files[0])
print(f"Testing file: {fp}")
print(f"Size: {os.path.getsize(fp)} bytes")

# openpyxl
wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
ws = wb.active
print(f"\nSheet: {ws.title}")
print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")

print("\n--- First 5 rows ---")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i >= 5:
        break
    vals = [str(c)[:80] if c is not None else None for c in row[:10]]
    print(f"  Row {i}: {vals}")

wb.close()

