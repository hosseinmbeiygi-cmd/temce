"""XLSX adapter — fast, memory-efficient parsing with openpyxl read_only mode."""

from __future__ import annotations

import logging
from pathlib import Path

from bulk_importer.adapters.base import BaseAdapter, ExtractedTable, ParseResult
from bulk_importer.config import MAX_ROWS_PER_TABLE, MAX_TABLES_PER_FILE

logger = logging.getLogger(__name__)


class XlsxAdapter(BaseAdapter):
    """Parse OOXML .xlsx files using openpyxl in read_only mode."""

    def parse(self, file_path: str) -> ParseResult:
        result = ParseResult(file_path=file_path, detected_format="xlsx_ooxml")

        try:
            from openpyxl import load_workbook
        except ImportError:
            result.errors.append("openpyxl not installed: pip install openpyxl")
            return result

        p = Path(file_path)
        if not p.exists():
            result.errors.append(f"File not found: {file_path}")
            return result

        try:
            wb = load_workbook(str(p), read_only=True, data_only=True)
        except Exception as e:
            result.errors.append(f"Failed to open workbook: {e}")
            return result

        try:
            for sheet_idx, ws in enumerate(wb.worksheets):
                if sheet_idx >= MAX_TABLES_PER_FILE:
                    logger.warning("Skipping sheets beyond %d in %s", MAX_TABLES_PER_FILE, file_path)
                    break

                headers: list[str] = []
                rows: list[list] = []
                row_count = 0

                for _row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                    if row_count >= MAX_ROWS_PER_TABLE:
                        break

                    # Convert all cells to strings for uniform handling
                    cells = [self._cell_to_str(c) for c in row]

                    # Skip completely empty rows
                    if not any(c for c in cells):
                        continue

                    # First non-empty row becomes headers
                    if not headers:
                        headers = cells
                        continue

                    rows.append(cells)
                    row_count += 1

                if headers or rows:
                    table = ExtractedTable(
                        table_index=sheet_idx,
                        sheet_name=ws.title,
                        headers=headers,
                        rows=rows,
                        row_count=row_count,
                        column_count=len(headers),
                    )
                    result.tables.append(table)
                    result.total_rows += row_count

        finally:
            wb.close()

        return result

    @staticmethod
    def _cell_to_str(value) -> str:
        """Convert a cell value to a clean string."""
        if value is None:
            return ""
        if isinstance(value, (int, float)):
            # Preserve numeric precision
            if isinstance(value, float) and value == int(value):
                return str(int(value))
            return str(value)
        return str(value).strip()
